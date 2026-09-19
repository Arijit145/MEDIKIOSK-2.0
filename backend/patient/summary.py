import json
from pathlib import Path

import pymupdf

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.enum.section import WD_SECTION


def add_prescription_to_document(document, prescription_path):

    if not prescription_path:
        document.add_paragraph(
            "No previous prescription was uploaded."
        )
        return

    file_path = Path(prescription_path)

    if not file_path.exists():
        document.add_paragraph(
            "Previous prescription file could not be found."
        )
        return

    extension = file_path.suffix.lower()

    # --------------------------------------------------
    # PDF PRESCRIPTION
    # --------------------------------------------------

    if extension == ".pdf":

        document.add_paragraph(
            "Original Prescription"
        ).runs[0].bold = True

        document.add_paragraph(
            "The original prescription is preserved separately. "
            "The pages below are included in this DOCX for doctor review."
        )

        pdf_document = pymupdf.open(str(file_path))

        for page_number, page in enumerate(
            pdf_document
        ):

            image = page.get_pixmap(
                matrix=pymupdf.Matrix(1.5, 1.5),
                alpha=False
            )

            image_path = (
                file_path.parent
                / f"_prescription_page_{page_number + 1}.png"
            )

            image.save(
                str(image_path)
            )

            document.add_paragraph(
                f"Prescription Page {page_number + 1}"
            )

            document.add_picture(
                str(image_path),
                width=Inches(6.0)
            )

            # Remove temporary image
            image_path.unlink(
                missing_ok=True
            )

        pdf_document.close()

    # --------------------------------------------------
    # IMAGE PRESCRIPTION
    # --------------------------------------------------

    elif extension in [
        ".jpg",
        ".jpeg",
        ".png"
    ]:

        document.add_paragraph(
            "Original Prescription"
        ).runs[0].bold = True

        document.add_paragraph(
            "The original prescription is preserved separately. "
            "The image below is included for doctor review."
        )

        document.add_picture(
            str(file_path),
            width=Inches(6.0)
        )

    # --------------------------------------------------
    # OTHER FILE TYPES
    # --------------------------------------------------

    else:

        document.add_paragraph(
            "Previous prescription uploaded:"
        )

        document.add_paragraph(
            file_path.name
        )

        document.add_paragraph(
            "The original prescription file is stored separately "
            "in its original format."
        )


def create_patient_summary(
    patient,
    symptom_record
):

    output_directory = Path(
        "uploads/summaries"
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True
    )

    file_name = (
        f"patient_summary_"
        f"{patient.patient_id.replace('/', '_')}.docx"
    )

    file_path = (
        output_directory / file_name
    )

    document = Document()

    # --------------------------------------------------
    # DOCUMENT SETTINGS
    # --------------------------------------------------

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        "MEDIKIOSK"
    )

    run.bold = True
    run.font.size = Pt(16)

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        "Patient Case Summary"
    )

    run.bold = True
    run.font.size = Pt(12)

    # --------------------------------------------------
    # PATIENT INFORMATION
    # --------------------------------------------------

    document.add_heading(
        "Patient Information",
        level=2
    )

    table = document.add_table(
        rows=5,
        cols=2
    )

    patient_information = [
        ("Patient ID", patient.patient_id),
        ("Name", patient.name),
        ("Age", str(patient.age)),
        ("Gender", patient.gender),
        ("Weight", f"{patient.weight} kg")
    ]

    for row, data in zip(
        table.rows,
        patient_information
    ):

        row.cells[0].text = data[0]
        row.cells[1].text = data[1]

        row.cells[0].paragraphs[0].runs[0].bold = True

    # --------------------------------------------------
    # REPORTED SYMPTOMS
    # --------------------------------------------------

    document.add_heading(
        "Reported Symptoms",
        level=2
    )

    document.add_paragraph(
        symptom_record.symptoms
    )

    # --------------------------------------------------
    # QUESTIONS AND ANSWERS
    # --------------------------------------------------

    document.add_heading(
        "History Questions and Answers",
        level=2
    )

    if symptom_record.answers:

        try:

            answers = json.loads(
                symptom_record.answers
            )

        except json.JSONDecodeError:

            answers = []

        questions = {}

        if symptom_record.generated_questions:

            try:

                generated_data = json.loads(
                    symptom_record.generated_questions
                )

                for question in generated_data.get(
                    "questions",
                    []
                ):

                    questions[
                        question.get("question_id")
                    ] = question.get("question")

            except json.JSONDecodeError:

                pass

        for answer in answers:

            question_id = answer.get(
                "question_id"
            )

            question_text = questions.get(
                question_id,
                question_id
            )

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                question_text + ": "
            )

            run.bold = True

            selected_options = answer.get(
                "selected_options",
                []
            )

            free_text = answer.get(
                "free_text"
            )

            answer_text = ", ".join(
                selected_options
            )

            if free_text:

                if answer_text:
                    answer_text += " - "

                answer_text += free_text

            paragraph.add_run(
                answer_text
                if answer_text
                else "Not provided"
            )

    else:

        document.add_paragraph(
            "No follow-up answers recorded."
        )

    # --------------------------------------------------
    # FAMILY HISTORY
    # --------------------------------------------------

    document.add_heading(
        "Family / Household History",
        level=2
    )

    if symptom_record.family_history:

        document.add_paragraph(
            symptom_record.family_history
        )

    else:

        document.add_paragraph(
            "No family or household history provided."
        )

    # --------------------------------------------------
    # PREVIOUS MEDICATION
    # --------------------------------------------------

    document.add_heading(
        "Previous Medication / Consultation",
        level=2
    )

    if symptom_record.previous_medication is True:

        document.add_paragraph(
            "Patient reported previous medication "
            "or medical consultation."
        )

    elif symptom_record.previous_medication is False:

        document.add_paragraph(
            "Patient reported no previous medication "
            "or medical consultation."
        )

    else:

        document.add_paragraph(
            "Not provided."
        )

    # --------------------------------------------------
    # PREVIOUS PRESCRIPTION
    # --------------------------------------------------

    document.add_heading(
        "Previous Prescription",
        level=2
    )

    add_prescription_to_document(
        document,
        symptom_record.prescription_path
    )

    # --------------------------------------------------
    # DOCTOR REVIEW
    # --------------------------------------------------

    document.add_heading(
        "Doctor Review",
        level=2
    )

    document.add_paragraph(
        "Doctor's notes / final assessment:"
    )

    document.add_paragraph(
        "____________________________________________"
    )

    document.add_paragraph(
        "____________________________________________"
    )

    document.add_paragraph(
        "____________________________________________"
    )

    # --------------------------------------------------
    # SAVE DOCUMENT
    # --------------------------------------------------

    document.save(
        file_path
    )

    return file_path